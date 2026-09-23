#!/usr/bin/env python3
"""Каталог оснасток из прайса в content/mounts.json.

    python3 scripts/parse_catalog.py ~/Downloads/osnastki_trodat_colop.xlsx

В книге по листу на серию, у всех одинаковая шапка. Размер в таблице — это
поле под клише: для круглых оснасток диаметр, для прямоугольных длина и
ширина. По нему и идёт подбор: печать Ø40 мм встанет только в оснастку с
полем Ø40 мм.
"""
from __future__ import annotations

import json
import re
import sys
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "content" / "mounts.json"

# лист -> бренд и семейство
SHEETS = {
    "Trodat Printy": ("Trodat", "Printy"),
    "IDEAL by Trodat": ("Trodat", "IDEAL"),
    "Trodat Professional": ("Trodat", "Professional"),
    "Trodat карманные": ("Trodat", "Карманные"),
    "Colop Printer": ("Colop", "Printer"),
    "Colop металл": ("Colop", "Металл"),
    "Colop карманные": ("Colop", "Карманные"),
}
HEADERS = ["Серия", "Модель", "Форма", "Тип"]


def num(value) -> float | None:
    if value is None:
        return None
    s = str(value).replace(",", ".").strip()
    m = re.search(r"\d+(?:\.\d+)?", s)
    return float(m.group(0)) if m else None


def main() -> int:
    src = Path(sys.argv[1] if len(sys.argv) > 1 else "~/Downloads/osnastki_trodat_colop.xlsx").expanduser()
    wb = load_workbook(src, data_only=True)

    items, skipped = [], 0
    for sheet, (brand, family) in SHEETS.items():
        if sheet not in wb.sheetnames:
            continue
        ws = wb[sheet]
        rows = list(ws.iter_rows(values_only=True))
        head_idx = next(i for i, r in enumerate(rows) if r and r[0] == "Серия")
        cols = {str(c).strip(): i for i, c in enumerate(rows[head_idx]) if c}

        def cell(row, name, default=None):
            for key, idx in cols.items():
                if key.startswith(name):
                    v = row[idx]
                    return v if v not in (None, "", "н/д") else default
            return default

        for row in rows[head_idx + 1:]:
            if not row or not row[cols["Модель"]]:
                continue
            price = num(cell(row, "Цена клиенту"))
            if price is None:          # без цены позиция бесполезна в каталоге
                skipped += 1
                continue
            d = num(cell(row, "Диаметр"))
            length, width = num(cell(row, "Длина")), num(cell(row, "Ширина"))
            shape = str(cell(row, "Форма", "")).strip()
            kind = str(cell(row, "Тип", "")).strip()
            items.append({
                "id": f"{brand}-{cell(row, 'Модель')}".replace(" ", "-").lower(),
                "brand": brand,
                "family": family,
                "series": str(cell(row, "Серия", family)),
                "model": str(cell(row, "Модель")),
                "shape": "round" if shape.startswith("Круг") else "rect",
                # тип в таблице пишется по-разному: «Карманная (самокрасящая)» и т.п.
                "kind": ("pocket" if kind.startswith("Карман")
                         else "auto" if kind.startswith("Авто")
                         else "manual"),
                "kindLabel": kind,
                "diameterMm": d,
                "lengthMm": length,
                "widthMm": width,
                "material": str(cell(row, "Материал", "")) or None,
                "colors": str(cell(row, "Цвета", "")) or None,
                "kit": str(cell(row, "Крышка", "")) or None,
                "price": int(price),
                "note": str(cell(row, "Примечание", "")) or None,
            })

    items.sort(key=lambda x: (x["shape"] != "round", x["price"]))
    sizes = sorted({i["diameterMm"] for i in items if i["diameterMm"]})
    OUT.write_text(json.dumps({"items": items, "roundSizes": sizes}, ensure_ascii=False, indent=1),
                   encoding="utf-8")

    print(f"позиций: {len(items)} (без цены пропущено {skipped}) → {OUT.relative_to(ROOT)}")
    print("круглые Ø:", ", ".join(f"{s:g}" for s in sizes))
    for shape in ("round", "rect"):
        part = [i for i in items if i["shape"] == shape]
        print(f"  {shape}: {len(part)}, цены {min(i['price'] for i in part)}–{max(i['price'] for i in part)} ₽")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
